#!/usr/bin/env python3
"""Rebuild Paper1_Extraction_Workbook.xlsx for EvalBot JAIR PRISMA review.

Standardizes all 4 sheets to align with:
  - config/data_dictionary.yaml (40 extraction fields)
  - config/inclusion_exclusion.yaml (screening criteria)
  - config/quality_appraisal.yaml (14-dimension rubric)
  - PRISMA 2020 flow (Identification -> Screening -> Included)

Usage:
    python scripts/build_extraction_workbook.py
"""
from __future__ import annotations

import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import OUTPUTS_DIR  # noqa: E402

WORKBOOK_PATH = Path(__file__).resolve().parent.parent / "workbook" / "Paper1_Extraction_Workbook.xlsx"

# Styles
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
SUBHEADER_FONT = Font(bold=True, size=10, color="2F5496")
SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
TITLE_FONT = Font(bold=True, size=14, color="2F5496")
SUBTITLE_FONT = Font(bold=True, size=11, color="2F5496")
BODY_FONT = Font(size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _header_row(ws, row, cols, widths=None):
    for c, val in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=val)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
        cell.border = THIN_BORDER
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def _body_row(ws, row, vals):
    for c, val in enumerate(vals, 1):
        cell = ws.cell(row=row, column=c, value=val)
        cell.font = BODY_FONT
        cell.alignment = WRAP
        cell.border = THIN_BORDER


def _section_header(ws, row, text, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = SUBTITLE_FONT
    cell.fill = SUBHEADER_FILL
    cell.alignment = WRAP


# ── Sheet 1: Instructions ─────────────────────────────────────────────

def _build_instructions(wb):
    ws = wb.create_sheet("Instructions")
    lines = [
        ("Paper 1 Literature Extraction Workbook", TITLE_FONT),
        ("", None),
        ("Project: EvalBot — Trustworthy Multimodal Generative AI for Automated Educational Assessment", SUBTITLE_FONT),
        ("Paper: JAIR PRISMA 2020 Systematic Review", SUBTITLE_FONT),
        ("", None),
        ("PURPOSE", SUBTITLE_FONT),
        ("Single place to log every candidate paper from the PRISMA search and extract structured data during full-text screening. All data entry is AI-draft until human-verified.", BODY_FONT),
        ("", None),
        ("HOW TO USE EACH SHEET", SUBTITLE_FONT),
        ("", None),
        ("1. Candidate Corpus — one row per paper.", SUBTITLE_FONT),
        ("   Fill bibliographic columns immediately. Fill extraction columns ONLY after reading the full text.", BODY_FONT),
        ("   screening_status tracks the PRISMA flow: Identified -> Abstract Screened -> Full Text Read -> Extracted -> Included in Synthesis.", BODY_FONT),
        ("   All extraction fields follow config/data_dictionary.yaml (40 fields). Missing values = 'NR' (never blank, never guessed).", BODY_FONT),
        ("", None),
        ("2. Gap Mapping — one row per research gap / verdict stage.", SUBTITLE_FONT),
        ("   Summarize what the screened corpus solves vs. what remains open, aligned to the 4-paper program.", BODY_FONT),
        ("   Includes a PRISMA 2020 flow table (Identification -> Screening -> Excluded -> Included).", BODY_FONT),
        ("", None),
        ("3. Quality Rubric — reference only.", SUBTITLE_FONT),
        ("   14-dimension quality appraisal rubric from config/quality_appraisal.yaml.", BODY_FONT),
        ("   Each dimension scored: strong | adequate | weak | NR (not reported).", BODY_FONT),
        ("   No composite score — each dimension reported individually.", BODY_FONT),
        ("", None),
        ("SCREENING WORKFLOW", SUBTITLE_FONT),
        ("   AI draft (scripts/08_screen_title_abstract.py) -> Human review -> title_abstract_human_decision", BODY_FONT),
        ("   Full-text: AI draft (prompts/07_full_text_reviewer.md) -> Human review -> fulltext_human_decision", BODY_FONT),
        ("   final_included = true ONLY after both human decisions are 'include'.", BODY_FONT),
        ("", None),
        ("EXTRACTION WORKFLOW", SUBTITLE_FONT),
        ("   Read PDF -> Apply prompts/08_evidence_extractor.md -> Populate extraction columns -> extraction_status = ai_draft", BODY_FONT),
        ("   Human verification -> extraction_status = human_verified", BODY_FONT),
        ("", None),
        ("QUALITY APPRAISAL WORKFLOW", SUBTITLE_FONT),
        ("   Read PDF + extraction data -> Apply prompts/09_quality_appraiser.md -> Rate 14 dimensions -> quality_appraisal_status = ai_draft", BODY_FONT),
        ("   Human verification -> quality_appraisal_status = human_verified", BODY_FONT),
        ("", None),
        ("PRISMA 2020 COMPLIANCE", SUBTITLE_FONT),
        ("   - Identification: all records from OpenAlex + Crossref (5 concept families each)", BODY_FONT),
        ("   - Screening: title/abstract (400 shortlisted) -> full-text (included papers)", BODY_FONT),
        ("   - Exclusion: every excluded record has an E-code (E01-E15) from config/inclusion_exclusion.yaml", BODY_FONT),
        ("   - Included: records with final_included = true -> evidence extraction -> quality appraisal", BODY_FONT),
        ("", None),
        ("MISSING VALUE CONVENTION", SUBTITLE_FONT),
        ("   'NR' = not reported / cannot be determined. Never use blank, N/A, or placeholder text.", BODY_FONT),
        ("", None),
        ("PROTOCOL STATUS", SUBTITLE_FONT),
        ("   FROZEN: review question, RQs, search string, inclusion/exclusion criteria, quality rubric, table designs.", BODY_FONT),
        ("   NOT FROZEN: corpus (keep adding rows as new papers appear, screened against frozen criteria).", BODY_FONT),
    ]
    for i, (text, font) in enumerate(lines, 1):
        cell = ws.cell(row=i, column=1, value=text)
        if font:
            cell.font = font
        cell.alignment = WRAP
    ws.column_dimensions["A"].width = 120


# ── Sheet 2: Candidate Corpus ─────────────────────────────────────────

CANDIDATE_HEADERS = [
    # Bibliographic (from master)
    "record_id",
    "title",
    "publication_year",
    "journal_name",
    "document_type",
    "doi_normalized",
    "doi_url",
    "authors",
    "source_database",
    # Screening
    "screening_status",
    "title_abstract_ai_decision",
    "title_abstract_human_decision",
    "fulltext_ai_decision",
    "fulltext_human_decision",
    "final_included",
    "exclusion_reason_code",
    # OA
    "open_access_status",
    "oa_best_status",
    "oa_evidence_source",
    # Relevance (from shortlist)
    "relevance_score",
    "matched_axes",
    # Educational context
    "educational_context",
    "educational_level",
    "subject_area",
    "assessment_task",
    "answer_format",
    # Architecture
    "input_modalities",
    "document_or_image_type",
    "ocr_or_document_model",
    "generative_model",
    "retrieval_or_rag_component",
    "rubric_or_reference_material",
    # Trustworthiness
    "evidence_grounding_method",
    "hallucination_definition_or_test",
    "hallucination_metric_result",
    "calibration_method",
    "calibration_metric_result",
    "uncertainty_or_abstention_method",
    "explainability_method",
    "fairness_or_bias_method",
    "human_in_the_loop_design",
    # Data & evaluation
    "dataset_name",
    "dataset_size",
    "annotation_and_ground_truth",
    "data_split_leakage_controls",
    "baselines",
    "primary_metrics",
    "key_results",
    # Reporting
    "main_findings",
    "author_reported_limitations",
    "ethics_consent_privacy",
    "funding_conflicts",
    "code_available",
    "data_available",
    # Quality appraisal (14 dimensions — populated by quality appraisal script)
    "qa_dataset_and_consent",
    "qa_ground_truth_construction",
    "qa_leakage_safeguards",
    "qa_baseline_comparisons",
    "qa_evaluation_metrics",
    "qa_calibration",
    "qa_grounding",
    "qa_hallucination_testing",
    "qa_fairness",
    "qa_robustness",
    "qa_human_evaluation",
    "qa_reproducibility",
    "qa_limitations_reporting",
    "qa_funding_conflicts",
    # Provenance
    "extraction_status",
    "quality_appraisal_status",
    "evidence_page_section",
    "citation_key",
    "notes",
]

CANDIDATE_WIDTHS = [
    12, 50, 8, 30, 12, 25, 35, 40, 12,  # bibliographic
    16, 18, 18, 18, 18, 12, 16,  # screening
    14, 12, 14,  # OA
    10, 20,  # relevance
    20, 14, 20, 25, 18,  # education
    25, 18, 20, 25, 20, 20,  # architecture
    25, 25, 20, 20, 20, 20, 20, 20, 20,  # trustworthiness
    25, 12, 25, 20, 25, 20, 25,  # data & eval
    40, 30, 25, 20, 10, 10,  # reporting
    12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12,  # QA dimensions
    14, 16, 20, 14, 30,  # provenance
]

CANDIDATE_DESCRIPTIONS = [
    # Row 2: descriptions for each column
    "FK to master_records_csv",
    "Paper title",
    "Publication year (2025-2026)",
    "Journal or conference name",
    "journal-article / proceedings-article",
    "Lowercase DOI (10.xxxx/xxxxx)",
    "https://doi.org/<doi_normalized>",
    "Semicolon-separated (Given Family)",
    "OpenAlex / Crossref",
    "Identified / Abstract Screened / Full Text Read / Extracted / Included in Synthesis",
    "include / exclude / unclear / pending",
    "include / exclude / unclear / pending",
    "include / exclude / unclear / pending",
    "include / exclude / unclear / pending",
    "true / false / pending",
    "E01-E15 (from inclusion_exclusion.yaml)",
    "gold / green / hybrid / bronze / closed / NR",
    "gold / hybrid / green / bronze / licensed / closed / unknown",
    "OpenAlex / Crossref / Unpaywall / unknown",
    "Relevance pre-filter score",
    "Matched relevance axes",
    "University / K-12 / vocational / mixed",
    "K-12 / higher-ed / vocational / mixed",
    "STEM / humanities / language / mixed",
    "short-answer / essay / diagram / mixed",
    "short-answer / essay / diagram / mixed",
    "semicolon-separated: image; OCR text; rubric; question",
    "handwritten / scanned / digital / diagram / mixed",
    "OCR engine or document understanding model",
    "e.g. GPT-4V, LLaVA, Gemini, Claude",
    "RAG pipeline description if used",
    "Rubric, reference answer, or exemplar used",
    "How generated scores are grounded to evidence",
    "Hallucination definition and test protocol",
    "Numeric result (e.g. hallucination rate = 12%)",
    "Calibration method (ECE, Brier, reliability diagram)",
    "Numeric result (e.g. ECE = 0.04)",
    "Selective prediction / abstention / deferral design",
    "Explainability method (attention, SHAP, rationale text)",
    "Fairness metric and subgroup breakdown",
    "Teacher review / override / human-in-the-loop interface",
    "Dataset name and source",
    "Number of students / scripts / items",
    "How ground truth was established (expert grading, IAA)",
    "Train/val/test split and leakage controls",
    "Baselines compared against",
    "Primary evaluation metrics",
    "Top-line numeric results",
    "2-4 sentence factual summary of findings",
    "Author-reported limitations",
    "Ethics approval, consent, privacy measures",
    "Funding sources and conflicts of interest",
    "yes / no / NR",
    "yes / no / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "strong / adequate / weak / NR",
    "ai_draft / human_verified / needs_review",
    "ai_draft / human_verified / needs_review",
    "Page/section/table reference for major claims",
    "e.g. Smith2025",
    "Free-text notes",
]


def _build_candidate_corpus(wb):
    ws = wb.create_sheet("Candidate Corpus")
    _header_row(ws, 1, CANDIDATE_HEADERS, CANDIDATE_WIDTHS)
    _body_row(ws, 2, CANDIDATE_DESCRIPTIONS)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CANDIDATE_HEADERS))}1"
    ws.freeze_panes = "A3"


# ── Sheet 3: Gap Mapping ──────────────────────────────────────────────

def _build_gap_mapping(wb):
    ws = wb.create_sheet("Gap Mapping")

    # Section 1: Research Gap Mapping
    _section_header(ws, 1, "Research Gap Mapping — EvalBot JAIR PRISMA Review", 6)
    ws.cell(row=2, column=1, value="Aligned to the 4-paper program. Update after full-text extraction.").font = BODY_FONT

    gap_headers = [
        "EvalBot Stage", "Papers Addressing It (of n=included)",
        "What They Solve", "Remaining Limitation (scope: included studies)",
        "Future Direction / Open Challenge", "Notes"
    ]
    _header_row(ws, 4, gap_headers, [30, 20, 35, 35, 35, 25])

    gap_rows = [
        ["Multimodal Assessment Architecture",
         "", "", "", "", "P1 scope: what architectures exist for multimodal educational assessment?"],
        ["Trustworthiness Evaluation (hallucination, calibration, grounding)",
         "", "", "", "", "P1 scope: how is trustworthiness measured in assessment systems?"],
        ["Generative Scoring & Feedback",
         "", "", "", "", "P3 scope: how do generative models produce scores/feedback?"],
        ["Human-AI Deferral & Oversight",
         "", "", "", "", "P4 scope: when should the model defer to human graders?"],
        ["Evidence Grounding & Attribution",
         "", "", "", "", "How are generated scores grounded to source evidence?"],
        ["Fairness & Bias Across Student Subgroups",
         "", "", "", "", "Bias across handwriting styles, language backgrounds, demographics?"],
        ["Reproducibility & Open Science",
         "", "", "", "", "Code/data availability, dataset public release, independent reproduction?"],
        ["Calibration & Uncertainty Quantification",
         "", "", "", "", "Is model confidence well-calibrated? Selective prediction supported?"],
    ]
    for i, row in enumerate(gap_rows, 5):
        _body_row(ws, i, row)

    # Section 2: PRISMA 2020 Flow
    prisma_start = 15
    _section_header(ws, prisma_start, "PRISMA 2020 Flow — Phase A-C Counts", 6)

    prisma_headers = ["Stage", "Count", "Source", "Notes", "", ""]
    _header_row(ws, prisma_start + 1, prisma_headers, [30, 12, 25, 40, 10, 10])

    prisma_rows = [
        ["Identification: records from databases", "", "OpenAlex (5 families) + Crossref (5 families)", "master_records.csv"],
        ["Identification: records after deduplication", "12028", "04_normalize_deduplicate.py", "Unique records (DOI-based)"],
        ["Screening: title/abstract screened", "400", "06c_relevance_oa_shortlist.py + 08_screen_title_abstract.py", "Relevance-scored + OA-prioritized"],
        ["Screening: excluded (title/abstract)", "", "E01-E15 codes", "See exclusion reason breakdown"],
        ["Screening: full-text assessed", "", "", "Pending human review"],
        ["Screening: excluded (full-text)", "", "E01-E15 codes", "Pending"],
        ["Included: final_included = true", "", "", "Evidence extraction + quality appraisal pending"],
    ]
    for i, row in enumerate(prisma_rows, prisma_start + 2):
        _body_row(ws, i, row)

    # Section 3: Pipeline Coverage Matrix
    coverage_start = prisma_start + 12
    _section_header(ws, coverage_start, "Pipeline Coverage Matrix — trustworthiness dimensions in included studies", 6)

    cov_headers = ["Trustworthiness Dimension", "Full (Y)", "Partial", "None / Not Stated", "Pending", "Total"]
    _header_row(ws, coverage_start + 1, cov_headers, [30, 10, 10, 15, 10, 10])

    cov_rows = [
        ["Hallucination / Factuality", "", "", "", "", ""],
        ["Calibration / Uncertainty", "", "", "", "", ""],
        ["Grounding / Attribution", "", "", "", "", ""],
        ["Explainability", "", "", "", "", ""],
        ["Fairness / Bias", "", "", "", "", ""],
        ["Robustness", "", "", "", "", ""],
        ["Human-in-the-loop / Abstention / Deferral", "", "", "", "", ""],
        ["Auditability", "", "", "", "", ""],
    ]
    for i, row in enumerate(cov_rows, coverage_start + 2):
        _body_row(ws, i, row)


# ── Sheet 4: Quality Rubric ──────────────────────────────────────────

def _build_quality_rubric(wb):
    ws = wb.create_sheet("Quality Rubric")

    _section_header(ws, 1, "14-Dimension Quality Appraisal Rubric — PRISMA 2020 Aligned", 5)
    ws.cell(row=2, column=1, value="Source: config/quality_appraisal.yaml | Each dimension scored independently. No composite score.").font = BODY_FONT

    rubric_headers = ["#", "Dimension", "Appraisal Question", "Scale", "Source in Extraction"]
    _header_row(ws, 4, rubric_headers, [4, 25, 60, 20, 25])

    dimensions = [
        ["1", "dataset_and_consent", "Is the dataset source described, and is participant/student consent or ethical approval documented?", "strong / adequate / weak / NR", "ethics_consent_privacy"],
        ["2", "ground_truth_construction", "Is the process for establishing ground-truth scores/labels described and reproducible?", "strong / adequate / weak / NR", "annotation_and_ground_truth"],
        ["3", "leakage_safeguards", "Are train/validation/test splits and measures against data leakage explicitly reported?", "strong / adequate / weak / NR", "data_split_leakage_controls"],
        ["4", "baseline_comparisons", "Are results compared against meaningful baselines (human graders, prior systems, ablations)?", "strong / adequate / weak / NR", "baselines"],
        ["5", "evaluation_metrics", "Are evaluation metrics appropriate, clearly defined, and reported with sufficient statistical detail?", "strong / adequate / weak / NR", "primary_metrics, key_results"],
        ["6", "calibration", "Is model confidence calibration evaluated (ECE, reliability diagrams, Brier score)?", "strong / adequate / weak / NR", "calibration_method, calibration_metric_result"],
        ["7", "grounding", "Is rubric/evidence grounding of generated scores or feedback explicitly evaluated?", "strong / adequate / weak / NR", "evidence_grounding_method"],
        ["8", "hallucination_testing", "Is hallucination or factuality explicitly defined and tested, with a reported metric?", "strong / adequate / weak / NR", "hallucination_definition_or_test, hallucination_metric_result"],
        ["9", "fairness", "Is fairness or bias across student subgroups evaluated?", "strong / adequate / weak / NR", "fairness_or_bias_method"],
        ["10", "robustness", "Is robustness tested (adversarial inputs, OCR noise, image quality variation)?", "strong / adequate / weak / NR", "QA checked from full text"],
        ["11", "human_evaluation", "Is there a human evaluation component beyond automated metrics?", "strong / adequate / weak / NR", "human_in_the_loop_design"],
        ["12", "reproducibility", "Are code, data, prompts, or model checkpoints available for independent reproduction?", "strong / adequate / weak / NR", "code_available, data_available"],
        ["13", "limitations_reporting", "Do the authors explicitly report limitations?", "strong / adequate / weak / NR", "author_reported_limitations"],
        ["14", "funding_conflicts", "Are funding sources and conflicts of interest disclosed?", "strong / adequate / weak / NR", "funding_conflicts"],
    ]
    for i, row in enumerate(dimensions, 5):
        _body_row(ws, i, row)

    # Scale definition
    scale_start = 20
    _section_header(ws, scale_start, "Scale Definitions", 5)
    scale_defs = [
        ["strong", "Explicit, rigorous treatment with verifiable detail"],
        ["adequate", "Present but limited detail or partial rigor"],
        ["weak", "Mentioned only superficially or with clear methodological gaps"],
        ["NR", "Not reported / cannot be determined from the text"],
    ]
    _header_row(ws, scale_start + 1, ["Rating", "Definition", "", "", ""], [12, 60, 10, 10, 10])
    for i, row in enumerate(scale_defs, scale_start + 2):
        _body_row(ws, i, row + ["", "", ""])

    # Policy
    policy_start = scale_start + 7
    _section_header(ws, policy_start, "Policy", 5)
    policies = [
        "No single composite score is computed automatically. Each dimension reported individually.",
        "Venue quartile, author reputation, and citation count must NOT influence quality ratings.",
        "Any narrative 'overall risk' judgment is AI-draft-only; requires explicit human sign-off before use in synthesis.",
        "Distinction: 'weak' = mentioned with clear gaps; 'NR' = not mentioned at all. Do not confuse the two.",
        "Every non-NR rating must include a one-line justification with page/section reference.",
    ]
    for i, p in enumerate(policies, policy_start + 1):
        cell = ws.cell(row=i, column=1, value=p)
        cell.font = BODY_FONT
        cell.alignment = WRAP


# ── Main ──────────────────────────────────────────────────────────────

def main():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    _build_instructions(wb)
    _build_candidate_corpus(wb)
    _build_gap_mapping(wb)
    _build_quality_rubric(wb)

    WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK_PATH)
    print(f"Workbook written: {WORKBOOK_PATH}")
    print(f"  Sheets: {wb.sheetnames}")
    print(f"  Candidate Corpus columns: {len(CANDIDATE_HEADERS)}")


if __name__ == "__main__":
    main()
